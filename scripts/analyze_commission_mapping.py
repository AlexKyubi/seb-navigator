from __future__ import annotations

import json
import re
from pathlib import Path

import openpyxl


ROOT = Path(__file__).resolve().parents[1]
CATALOG = ROOT / "Data" / "справочникSEB.xlsx"
FOCUS = ROOT / "Data" / "Фокусы Seb Сентябрь 2026 ит-1.xlsx"


def norm(value: object) -> str:
    return re.sub(r"[^A-Z0-9А-ЯЁ]", "", str(value or "").upper())


catalog_wb = openpyxl.load_workbook(CATALOG, read_only=True, data_only=True)
catalog_ws = catalog_wb.active
headers = {str(cell.value).strip(): cell.column for cell in catalog_ws[1] if cell.value is not None}

catalog = []
for values in catalog_ws.iter_rows(min_row=2, values_only=True):
    record = {name: values[column - 1] for name, column in headers.items()}
    record["_search"] = norm(" ".join(str(record.get(key) or "") for key in ("Comm.Code", "Наименование", "Sulpak.title")))
    record["_comm"] = norm(record.get("Comm.Code"))
    catalog.append(record)

focus_wb = openpyxl.load_workbook(FOCUS, read_only=True, data_only=True)
commission_rows = []
for ws in focus_wb.worksheets:
    header_row = None
    header_map = {}
    all_rows = list(ws.iter_rows(values_only=True))
    for row, values in enumerate(all_rows[:20], 1):
        if "Comm.Code" in values:
            header_row = row
            header_map = {" ".join(str(value).split()): index + 1 for index, value in enumerate(values) if value is not None}
            break
    if not header_row:
        continue
    comm_col = header_map["Comm.Code"]
    below_col = header_map.get("% Выполнение ниже плана")
    above_col = header_map.get("Выполнение выше плана") or header_map.get("% Выполнение выше плана")
    for row, values in enumerate(all_rows[header_row:], header_row + 1):
        code = values[comm_col - 1] if comm_col <= len(values) else None
        if not code:
            continue
        below = values[below_col - 1] if below_col and below_col <= len(values) else None
        above = values[above_col - 1] if above_col and above_col <= len(values) else None
        if not isinstance(below, (int, float)) or not isinstance(above, (int, float)):
            continue
        commission_rows.append({"sheet": ws.title, "row": row, "code": str(code).strip(), "below": below, "above": above})

results = []
for item in commission_rows:
    code = norm(item["code"])
    exact = [record for record in catalog if record["_comm"] and record["_comm"] == code]
    candidates = exact or [record for record in catalog if len(code) >= 5 and code in record["_search"]]
    results.append(
        {
            **item,
            "normalized": code,
            "match_type": "exact_comm" if exact else ("title_token" if candidates else "none"),
            "articles": [str(record.get("Артикул")) for record in candidates],
            "titles": [record.get("Sulpak.title") or record.get("Наименование") for record in candidates],
        }
    )

summary = {
    "commission_rows": len(results),
    "matched_unique": sum(len(item["articles"]) == 1 for item in results),
    "unmatched": sum(not item["articles"] for item in results),
    "ambiguous": sum(len(item["articles"]) > 1 for item in results),
    "matched_catalog_articles": len({article for item in results for article in item["articles"] if len(item["articles"]) == 1}),
    "unmatched_rows": [item for item in results if not item["articles"]],
    "ambiguous_rows": [item for item in results if len(item["articles"]) > 1],
}
print(json.dumps(summary, ensure_ascii=False, indent=2))

from __future__ import annotations

import argparse
import concurrent.futures
import copy
import datetime as dt
import json
import os
import time
import urllib.error
import urllib.request
from pathlib import Path
from typing import Any

import openpyxl


API_TEMPLATE = "https://sulpak-api.evinent.site/api/search/autocomplete/1/3/{article}/true/"
DEFAULT_FIELDS = [
    "code",
    "title",
    "brand",
    "url",
    "photoUrl",
    "price",
    "priceOld",
    "isAvailable",
    "properties",
]


def normalize_article(value: Any) -> str:
    if value is None:
        return ""
    if isinstance(value, float) and value.is_integer():
        return str(int(value))
    return str(value).strip()


def fetch_one(article: str, attempts: int = 4) -> dict[str, Any]:
    last_error = "unknown error"
    for attempt in range(1, attempts + 1):
        try:
            request = urllib.request.Request(
                API_TEMPLATE.format(article=article),
                headers={
                    "Accept": "application/json",
                    "Referer": "https://www.sulpak.kz/",
                    "User-Agent": "SebBN catalog updater/1.0",
                },
            )
            with urllib.request.urlopen(request, timeout=30) as response:
                payload = json.loads(response.read().decode("utf-8"))
                if not payload:
                    return {"status": "stale", "article": article, "reason": "empty_json"}
                products = payload.get("products") or []
                exact = next(
                    (item for item in products if normalize_article(item.get("code")) == article),
                    None,
                )
                if exact is None:
                    return {
                        "status": "stale",
                        "article": article,
                        "reason": "no_exact_product",
                    }
                return {"status": "found", "article": article, "product": exact}
        except urllib.error.HTTPError as exc:
            last_error = f"HTTP {exc.code}"
            if exc.code not in {408, 425, 429, 500, 502, 503, 504}:
                return {"status": "error", "article": article, "error": last_error}
        except (urllib.error.URLError, TimeoutError, ValueError) as exc:
            last_error = f"{type(exc).__name__}: {exc}"
        if attempt < attempts:
            time.sleep(0.7 * (2 ** (attempt - 1)))
    return {"status": "error", "article": article, "error": last_error}


def cell_value(value: Any) -> Any:
    if isinstance(value, (dict, list)):
        return json.dumps(value, ensure_ascii=False, separators=(",", ":"))
    return value


def copy_row_style(ws: openpyxl.worksheet.worksheet.Worksheet, source_row: int, target_row: int, max_col: int) -> None:
    ws.row_dimensions[target_row].height = ws.row_dimensions[source_row].height
    for col in range(1, max_col + 1):
        source = ws.cell(source_row, col)
        target = ws.cell(target_row, col)
        if source.has_style:
            target._style = copy.copy(source._style)
        if source.number_format:
            target.number_format = source.number_format
        if source.alignment:
            target.alignment = copy.copy(source.alignment)
        if source.protection:
            target.protection = copy.copy(source.protection)


def main() -> int:
    parser = argparse.ArgumentParser(description="Enrich the SEB product catalog from Sulpak autocomplete.")
    parser.add_argument("workbook", type=Path)
    parser.add_argument("--output", type=Path)
    parser.add_argument("--workers", type=int, default=6)
    args = parser.parse_args()

    workbook_path = args.workbook.resolve()
    output_path = (args.output or workbook_path.with_name(f"{workbook_path.stem}_обновленный.xlsx")).resolve()
    workbook = openpyxl.load_workbook(workbook_path, keep_vba=False, data_only=False)
    worksheet = workbook.worksheets[0]

    rows_by_article: dict[str, list[int]] = {}
    for row in range(2, worksheet.max_row + 1):
        article = normalize_article(worksheet.cell(row, 3).value)
        if article:
            rows_by_article.setdefault(article, []).append(row)

    articles = list(rows_by_article)
    results: dict[str, dict[str, Any]] = {}
    with concurrent.futures.ThreadPoolExecutor(max_workers=max(1, min(args.workers, 8))) as executor:
        futures = {executor.submit(fetch_one, article): article for article in articles}
        completed = 0
        for future in concurrent.futures.as_completed(futures):
            article = futures[future]
            try:
                results[article] = future.result()
            except Exception as exc:  # Defensive: retain the row on unexpected worker failure.
                results[article] = {"status": "error", "article": article, "error": repr(exc)}
            completed += 1
            if completed % 100 == 0 or completed == len(futures):
                print(f"Fetched {completed}/{len(futures)}", flush=True)

    found = {article: result["product"] for article, result in results.items() if result["status"] == "found"}
    stale = {article for article, result in results.items() if result["status"] == "stale"}
    errors = {article: result.get("error", "unknown") for article, result in results.items() if result["status"] == "error"}

    field_names = list(DEFAULT_FIELDS)
    for product in found.values():
        for key in product:
            if key not in field_names:
                field_names.append(key)

    existing_headers = {str(worksheet.cell(1, col).value): col for col in range(1, worksheet.max_column + 1)}
    field_columns: dict[str, int] = {}
    next_column = worksheet.max_column + 1
    for field in field_names:
        header = f"Sulpak.{field}"
        column = existing_headers.get(header)
        if column is None:
            column = next_column
            next_column += 1
            worksheet.cell(1, column).value = header
            source_header = worksheet.cell(1, min(worksheet.max_column, 5))
            if source_header.has_style:
                worksheet.cell(1, column)._style = copy.copy(source_header._style)
            worksheet.column_dimensions[openpyxl.utils.get_column_letter(column)].width = 18
        field_columns[field] = column

    max_output_col = max(field_columns.values(), default=worksheet.max_column)
    kept_rows: list[list[Any]] = []
    kept_styles: list[int] = []
    for row in range(2, worksheet.max_row + 1):
        article = normalize_article(worksheet.cell(row, 3).value)
        if article in stale:
            continue
        values = [worksheet.cell(row, col).value for col in range(1, max_output_col + 1)]
        if article in found:
            product = found[article]
            for field, column in field_columns.items():
                values[column - 1] = cell_value(product.get(field))
        kept_rows.append(values)
        kept_styles.append(row)

    original_last_row = worksheet.max_row
    for target_row, (values, source_row) in enumerate(zip(kept_rows, kept_styles), start=2):
        copy_row_style(worksheet, source_row, target_row, max_output_col)
        for column, value in enumerate(values, start=1):
            worksheet.cell(target_row, column).value = value
    new_last_row = len(kept_rows) + 1
    if original_last_row > new_last_row:
        worksheet.delete_rows(new_last_row + 1, original_last_row - new_last_row)

    if worksheet.auto_filter.ref:
        worksheet.auto_filter.ref = f"A1:{openpyxl.utils.get_column_letter(max_output_col)}{new_last_row}"
    for row in range(1, new_last_row + 1):
        worksheet.row_dimensions[row].hidden = False
    workbook.calculation.fullCalcOnLoad = True
    workbook.calculation.forceFullCalc = True
    for extra_sheet in workbook.worksheets[1:]:
        workbook.remove(extra_sheet)

    timestamp = dt.datetime.now().strftime("%Y%m%d-%H%M%S")
    report_path = output_path.with_name(f"sulpak-refresh-{timestamp}.json")
    temp_path = output_path.with_name(f".{output_path.stem}.tmp{output_path.suffix}")
    output_path.parent.mkdir(parents=True, exist_ok=True)
    workbook.save(temp_path)
    os.replace(temp_path, output_path)

    report = {
        "source_workbook": str(workbook_path),
        "output_workbook": str(output_path),
        "sheet": worksheet.title,
        "article_column": 3,
        "queried_unique_articles": len(articles),
        "found_unique_articles": len(found),
        "removed_unique_articles": len(stale),
        "retained_on_error_unique_articles": len(errors),
        "rows_before": original_last_row - 1,
        "rows_after": new_last_row - 1,
        "removed_articles": sorted(stale),
        "errors": errors,
        "metadata_columns": [f"Sulpak.{name}" for name in field_names],
    }
    report_path.write_text(json.dumps(report, ensure_ascii=False, indent=2), encoding="utf-8")
    print(json.dumps(report, ensure_ascii=False), flush=True)
    return 0 if not errors else 2


if __name__ == "__main__":
    raise SystemExit(main())

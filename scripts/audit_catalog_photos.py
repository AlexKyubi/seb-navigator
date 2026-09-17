from __future__ import annotations

import collections
import json
import sys
import urllib.parse
from pathlib import Path

import openpyxl


path = Path(sys.argv[1] if len(sys.argv) > 1 else "Data/справочникSEB.xlsx")
workbook = openpyxl.load_workbook(path, read_only=True, data_only=False)
sheet = workbook.active
headers = {str(cell.value).strip(): cell.column for cell in sheet[1] if cell.value is not None}
photo_column = headers.get("Sulpak.photoUrl")
article_column = headers.get("Артикул")
if not photo_column or not article_column:
    raise SystemExit("Required catalog columns are missing")

records = []
for row in range(2, sheet.max_row + 1):
    article = str(sheet.cell(row, article_column).value or "").strip()
    photo = str(sheet.cell(row, photo_column).value or "").strip()
    if article:
        records.append({"row": row, "article": article, "photo": photo})

domains = collections.Counter(urllib.parse.urlparse(item["photo"]).netloc for item in records if item["photo"])
url_counts = collections.Counter(item["photo"] for item in records if item["photo"])
result = {
    "workbook": str(path),
    "sheet": sheet.title,
    "rows": len(records),
    "empty_count": sum(not item["photo"] for item in records),
    "empty_sample": [item for item in records if not item["photo"]][:20],
    "invalid_format_count": sum(bool(item["photo"]) and not item["photo"].startswith(("https://", "http://")) for item in records),
    "invalid_format_sample": [item for item in records if item["photo"] and not item["photo"].startswith(("https://", "http://"))][:20],
    "duplicate_url_count": sum(count - 1 for count in url_counts.values() if count > 1),
    "duplicate_url_sample": dict(list({url: count for url, count in url_counts.items() if count > 1}.items())[:20]),
    "domains": domains.most_common(20),
}
payload = json.dumps(result, ensure_ascii=False, indent=2)
if len(sys.argv) > 2:
    output = Path(sys.argv[2])
    output.parent.mkdir(parents=True, exist_ok=True)
    output.write_text(payload, encoding="utf-8")
print(payload)
